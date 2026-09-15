from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib import colors

def exportar_asistencias_excel(asistencias):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Asistencias"
    encabezados = [
        "Socio",
        "Actividad",
        "Día",
        "Horario",
        "Profesor",
        "Fecha",
        "Hora",
        "Estado",
    ]
    hoja.append(encabezados)

    for asistencia in asistencias.iterator():
        hoja.append([
            str(asistencia.socio),
            asistencia.horario.actividad.nombre,
            asistencia.horario.dia,
            f"{asistencia.horario.hora_inicio:%H:%M} - "
            f"{asistencia.horario.hora_fin:%H:%M}",
            str(asistencia.profesor),
            asistencia.fecha,
            asistencia.hora,
            "Presente" if asistencia.presente else "Ausente",
        ])

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="041A30")

    anchos = [28, 24, 14, 20, 28, 14, 12, 14]
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[chr(64 + indice)].width = ancho

    for fila in range(2, hoja.max_row + 1):
        hoja.cell(fila, 6).number_format = "dd/mm/yyyy"
        hoja.cell(fila, 7).number_format = "HH:mm"

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    archivo = BytesIO()
    libro.save(archivo)
    respuesta = HttpResponse(
        archivo.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    respuesta["Content-Disposition"] = (
        'attachment; filename="asistencias.xlsx"'
    )
    return respuesta


def exportar_asistencias_pdf(asistencias):
    archivo = BytesIO()
    pdf = canvas.Canvas(archivo, pagesize=landscape(letter))
    ancho, alto = landscape(letter)
    margen = 15 * mm
    y = alto - margen

    def dibujar_encabezado():
        nonlocal y
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(margen, y, "Reporte de asistencias")
        y -= 10 * mm
        pdf.setFillColorRGB(4 / 255, 26 / 255, 48 / 255)
        pdf.rect(margen, y - 5 * mm, ancho - 2 * margen, 7 * mm, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 8)
        posiciones = [margen, 65 * mm, 105 * mm, 135 * mm, 170 * mm, 220 * mm, 245 * mm, 260 * mm]
        for posicion, titulo in zip(
            posiciones,
            ["Socio", "Actividad", "Día", "Horario", "Profesor", "Fecha", "Hora", "Estado"],
        ):
            pdf.drawString(posicion, y - 1 * mm, titulo)
        y -= 10 * mm

    dibujar_encabezado()
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica", 7)

    for asistencia in asistencias.iterator():
        if y < margen:
            pdf.showPage()
            y = alto - margen
            dibujar_encabezado()
            pdf.setFillColor(colors.black)
            pdf.setFont("Helvetica", 7)

        datos = [
            str(asistencia.socio)[:30],
            asistencia.horario.actividad.nombre[:24],
            asistencia.horario.dia[:14],
            f"{asistencia.horario.hora_inicio:%H:%M}-"
            f"{asistencia.horario.hora_fin:%H:%M}",
            str(asistencia.profesor)[:24],
            asistencia.fecha.strftime("%d/%m/%Y"),
            asistencia.hora.strftime("%H:%M"),
            "Presente" if asistencia.presente else "Ausente",
        ]
        posiciones = [margen, 65 * mm, 105 * mm, 135 * mm, 170 * mm, 220 * mm, 245 * mm, 260 * mm]
        for posicion, dato in zip(posiciones, datos):
            pdf.drawString(posicion, y, dato)
        y -= 6 * mm

    pdf.save()
    respuesta = HttpResponse(archivo.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = 'attachment; filename="asistencias.pdf"'
    return respuesta
