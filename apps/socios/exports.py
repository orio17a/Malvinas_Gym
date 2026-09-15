from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import openpyxl

def exportar_socios(socios):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Socios"
    hoja.append(["DNI", "Apellido", "Nombre", "Email", "Teléfono", "Estado", "Apto físico", "Ingreso"])
    for socio in socios.iterator():
        hoja.append([
            socio.dni, socio.apellido, socio.nombre, socio.email or "-",
            socio.telefono or "-", socio.estado_socio.nombre,
            socio.estado_apto_fisico.nombre, socio.fecha_ingreso,
        ])
        # Los datos del usuario siempre son texto, nunca fórmulas de Excel.
        for celda in hoja[hoja.max_row][:7]:
            celda.data_type = "s"
        hoja.cell(hoja.max_row, 8).number_format = "dd/mm/yyyy"
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="041A30")
    for columna, ancho in zip("ABCDEFGH", [18, 24, 24, 36, 20, 18, 20, 16]):
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    archivo = BytesIO()
    libro.save(archivo)
    respuesta = HttpResponse(archivo.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = 'attachment; filename="socios.xlsx"'
    return respuesta
